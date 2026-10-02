#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate files/cv.pdf and files/achievements.xlsx from _data/cv_data.yml.

The PDF replicates the layout of the original Word CV template (A4, SimSun
28 name, KaiTi body with Times New Roman for Latin text, DengXian-Bold
section headings, numbered entries with hanging indent, centered "-N-"
page footer). The Excel achievements list follows the same design: KaiTi
body, SimSun title, DengXian bold headers, black text on white.

Automation: papers newly added to _publications/ (matched by DOI, or by
title when no DOI is present) that are not yet cited in cv_data.yml are
auto-inserted into the publication list - author list and journal are
looked up from CrossRef when possible - and both files are regenerated.

Usage:
    python auto_cv.py            # regenerate files/cv.pdf + achievements.xlsx
    python auto_cv.py --dry-run  # only report, write nothing

Dependencies: reportlab, pyyaml, openpyxl
"""

import argparse
import html
import json
import re
import sys
import urllib.request
from pathlib import Path

import yaml
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import BaseDocTemplate, Frame, PageBreak, PageTemplate, Paragraph, Spacer

ROOT = Path(__file__).resolve().parent
DATA_FILE = ROOT / "_data" / "cv_data.yml"
PUBS_DIR = ROOT / "_publications"
OUT_FILE = ROOT / "files" / "cv.pdf"
XLSX_FILE = ROOT / "files" / "achievements.xlsx"
PHOTO_FILE = ROOT / "images" / "cv-photo.jpg"
PHOTO_W, PHOTO_H = 100, 133  # pt, 3:4

CROSSREF_API = "https://api.crossref.org/works"
CONTACT_EMAIL = "guoxiang@henu.edu.cn"
USER_AGENT = "GuoXiang.github.io CV builder (mailto:%s)" % CONTACT_EMAIL

# ---- geometry (pt), measured from the original template ----
MARGIN_L, MARGIN_R, MARGIN_T, MARGIN_B = 54, 40, 64, 64
CONTENT_W = A4[0] - MARGIN_L - MARGIN_R  # ~501pt

FONT_FILES = [
    ("KaiTi", r"C:\Windows\Fonts\simkai.ttf", 0),
    ("SimSun", r"C:\Windows\Fonts\simsun.ttc", 0),
    ("Heading", r"C:\Windows\Fonts\Dengb.ttf", 0),   # 等线 Bold, template heading font
    ("HeadingAlt", r"C:\Windows\Fonts\simhei.ttf", 0),
    ("Times", r"C:\Windows\Fonts\times.ttf", 0),
]


def register_fonts():
    from reportlab.pdfbase.pdfmetrics import registerFontFamily
    ok = {}
    for name, path, idx in FONT_FILES:
        if Path(path).exists():
            pdfmetrics.registerFont(TTFont(name, path, subfontIndex=idx))
            ok[name] = name
        else:
            ok[name] = None
    if not ok["Heading"]:
        ok["Heading"] = ok.get("HeadingAlt") or ok["KaiTi"]
    if not ok["Times"]:
        ok["Times"] = ok["KaiTi"]
    registerFontFamily("KaiTi", normal="KaiTi", bold="Heading", italic="KaiTi", boldItalic="Heading")
    return ok


LATIN = re.compile(r"[ -~]")


def has_alnum(s):
    return any(c.isalnum() for c in s)


def mixed(text, base="KaiTi", latin="Times"):
    """Escape and wrap Latin runs in <font> tags so Latin text uses Times."""
    text = html.escape(text, quote=False)
    runs, cur, cur_lat = [], "", None
    for ch in text:
        is_lat = bool(LATIN.match(ch))
        if cur_lat is None or is_lat == cur_lat:
            cur += ch
        else:
            runs.append((cur_lat, cur))
            cur = ch
        cur_lat = is_lat
    if cur:
        runs.append((cur_lat, cur))
    out = []
    for is_lat, run in runs:
        if is_lat and has_alnum(run) and latin != base:
            out.append('<font name="%s">%s</font>' % (latin, run))
        else:
            out.append(run)
    return "".join(out)


def styles():
    return {
        "name": ParagraphStyle("name", fontName="SimSun", fontSize=28, leading=34,
                               alignment=0, spaceAfter=0),
        "info": ParagraphStyle("info", fontName="KaiTi", fontSize=15, leading=23,
                               wordWrap="CJK"),
        "heading": ParagraphStyle("heading", fontName="Heading", fontSize=14, leading=17,
                                  spaceBefore=14, spaceAfter=2, wordWrap="CJK"),
        "bio": ParagraphStyle("bio", fontName="KaiTi", fontSize=14, leading=20,
                              firstLineIndent=28, wordWrap="CJK"),
        "entry": ParagraphStyle("entry", fontName="KaiTi", fontSize=12, leading=15,
                                leftIndent=5, wordWrap="CJK"),
        "sub": ParagraphStyle("sub", fontName="KaiTi", fontSize=12, leading=15,
                              leftIndent=17, wordWrap="CJK"),
        "numbered": ParagraphStyle("numbered", fontName="KaiTi", fontSize=12, leading=15,
                                   leftIndent=24, firstLineIndent=-24, wordWrap="CJK"),
    }


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("SimSun", 9)
    canvas.drawCentredString(A4[0] / 2, 25, "-%d-" % doc.page)
    canvas.restoreState()


def parse_front_matter(path):
    text = path.read_text(encoding="utf-8-sig").lstrip("\ufeff")
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n", text, re.S)
    if not m:
        return {}
    fm = {}
    for line in m.group(1).splitlines():
        mm = re.match(r"^([A-Za-z_]+):\s*['\"]?(.*?)['\"]?\s*$", line)
        if mm:
            fm[mm.group(1)] = mm.group(2)
    return fm


def doi_from_url(url):
    m = re.search(r"(10\.\d{4,9}/[^\s'\"]+?)\.pdf$", url or "")
    if m:
        return m.group(1)
    m = re.search(r"(10\.\d{4,9}/[^\s'\"]+)$", url or "")
    return m.group(1).rstrip(").,;") if m else ""


def norm_title(t):
    return re.sub(r"[^a-z0-9]+", "", (t or "").lower())


def crossref(doi):
    url = "%s/%s" % (CROSSREF_API, doi)
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=30) as r:
        m = json.load(r)["message"]
    year = (m.get("issued", {}).get("date-parts") or [[None]])[0][0]
    vol = m.get("volume", "")
    page = m.get("page", "") or m.get("article-number", "")
    vp = "; ".join(x for x in [str(year), vol + (": " + page if page and vol else (": " + page if page else ""))] if x)
    return {
        "authors": [((a.get("given", "") + " " + a.get("family", "")).strip()) for a in m.get("author", [])],
        "title": m.get("title", [""])[0],
        "journal": (m.get("container-title") or [""])[0].replace("&amp;", "&"),
        "year": year,
        "vp": vp,
        "doi": doi,
    }


CN_NUM = "一二三四五六七八九"


def author_note(authors, me="Xiang Guo"):
    if not authors:
        return "新收录"
    pos = next((i for i, a in enumerate(authors, 1) if me.split()[-1].lower() in a.lower()), 0)
    if pos == 1:
        return "第一作者完成"
    if pos == 2:
        return "第二作者完成"
    if pos == 3:
        return "第三作者完成"
    return "其他作者完成"


# Sections of the CV that are career history rather than research outputs;
# they are excluded from the Excel achievements list.
XLSX_EXCLUDE_SECTIONS = {"工作经历", "教育经历", "访问经历", "研究领域"}

XLSX_HEADERS = ["序号", "类别", "作者贡献", "成果内容", "年份", "DOI / 链接"]


def _entry_rows(sec, pubs):
    """Yield (no, note, text, year, doi) for one CV section."""
    if sec.get("publications"):
        for i, p in enumerate(pubs, 1):
            yield i, p.get("note", ""), p["text"], p.get("year") or "", p.get("doi", "")
        return
    for i, entry in enumerate(sec["entries"], 1):
        text = entry["text"]
        m = re.match(r"^\(([^)]*)\)\s*", text)
        note = m.group(1) if m else ""
        if m:
            text = text[m.end():]
        ym = re.search(r"(?<!\d)(?:19|20)\d{2}(?!\d)", text)
        yield i, note, text, int(ym.group()) if ym else "", ""


def build_xlsx(data, out_file):
    """Achievements list in Excel, styled after the CV template:
    KaiTi body, SimSun centered title, DengXian bold headers, black on
    white, numbered entries grouped in CV section order."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "成果列表"
    ws.sheet_view.showGridLines = False

    f_title = Font(name="SimSun", size=16, bold=True)
    f_head = Font(name="等线", size=12, bold=True)
    f_body = Font(name="KaiTi", size=12)
    f_link = Font(name="KaiTi", size=12, underline="single")
    thin_gray = Side(style="thin", color="BFBFBF")
    black = Side(style="thin", color="000000")
    row_border = Border(bottom=thin_gray)
    head_border = Border(bottom=black)

    # title (CV style: SimSun, centered)
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = data["personal"]["name"].replace(" ", "") + " 成果列表"
    c.font = f_title
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 6

    # header row
    for col, h in enumerate(XLSX_HEADERS, 1):
        c = ws.cell(row=3, column=col, value=h)
        c.font = f_head
        c.border = head_border
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 22

    pubs = data.get("publications", [])
    row = 4
    for sec in data["sections"]:
        if sec["title"] in XLSX_EXCLUDE_SECTIONS:
            continue
        for no, note, text, year, doi in _entry_rows(sec, pubs):
            values = [no, sec["title"], note, text, year or "",
                      ("https://doi.org/" + doi) if doi else ""]
            for col, v in enumerate(values, 1):
                c = ws.cell(row=row, column=col, value=v)
                c.font = f_body
                c.border = row_border
                if col in (1, 2, 5):
                    c.alignment = Alignment(horizontal="center", vertical="top")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            if doi:
                link = ws.cell(row=row, column=6)
                link.hyperlink = link.value
                link.font = f_link
            row += 1

    widths = [6, 14, 24, 95, 8, 34]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = "A4"
    if row > 4:
        ws.auto_filter.ref = "A3:F%d" % (row - 1)

    out_file.parent.mkdir(exist_ok=True)
    wb.save(out_file)
    return row - 4


def build_story(data, st):
    from reportlab.platypus import Image as RLImage
    from reportlab.platypus import Table, TableStyle

    story = []
    # ---- page 1: name (left) / photo (right) / info / bio ----
    left_flow = [Paragraph(data["personal"]["name"], st["name"]), Spacer(1, 33)]
    left_flow += [Paragraph(mixed(line), st["info"]) for line in data["personal"]["info_lines"]]

    if PHOTO_FILE.exists():
        photo = RLImage(str(PHOTO_FILE), width=PHOTO_W, height=PHOTO_H)
        header = Table([[left_flow, photo]],
                       colWidths=[CONTENT_W - PHOTO_W - 14, PHOTO_W + 14])
        header.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ALIGN", (1, 0), (1, 0), "RIGHT"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ]))
    else:
        header = left_flow

    story.append(Spacer(1, 26))
    if isinstance(header, Table):
        story.append(header)
    else:
        story.extend(header)
    story.append(Paragraph("●  个人简介", st["heading"]))
    story.append(Spacer(1, 4))
    for para in data["bio"]:
        story.append(Paragraph(mixed(para), st["bio"]))
    story.append(PageBreak())

    pubs = data["publications"]
    for sec in data["sections"]:
        story.append(Paragraph("●  %s" % sec["title"], st["heading"]))
        story.append(Spacer(1, 4))
        if sec.get("publications"):
            for i, p in enumerate(pubs, 1):
                text = "(%s) %s" % (p["note"], p["text"]) if p.get("note") else p["text"]
                story.append(Paragraph(mixed("%d. %s" % (i, text)), st["numbered"]))
            continue
        n = 0
        for entry in sec["entries"]:
            n += 1
            text = entry["text"]
            if sec.get("numbered"):
                story.append(Paragraph(mixed("%d. %s" % (n, text)), st["numbered"]))
            else:
                story.append(Paragraph(mixed(text), st["entry"]))
            for sub in entry.get("subs", []):
                story.append(Paragraph(mixed(sub), st["sub"]))
    return story


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    data = yaml.safe_load(DATA_FILE.read_text(encoding="utf-8"))
    pubs = data.setdefault("publications", [])

    known_dois = {p.get("doi", "").lower() for p in pubs if p.get("doi")}
    known_titles = [norm_title(p["text"]) for p in pubs]

    def known(md_title):
        t = norm_title(md_title)
        return any(t and t in kt for kt in known_titles)

    new_entries = []
    if PUBS_DIR.exists():
        for md in sorted(PUBS_DIR.glob("*.md")):
            fm = parse_front_matter(md)
            doi = doi_from_url(fm.get("paperurl", ""))
            if doi and doi.lower() in known_dois:
                continue
            if not doi and known(fm.get("title", "")):
                continue
            meta = None
            if doi:
                try:
                    meta = crossref(doi)
                except Exception as e:
                    print("  ! CrossRef lookup failed for %s: %s" % (doi, e))
            if meta is None:
                meta = {"authors": [], "title": fm.get("title", ""),
                        "journal": fm.get("venue", ""), "year": (fm.get("date", "") or "")[:4] or 0,
                        "vp": "", "doi": doi}
            if known(meta["title"]):
                continue
            authors = ", ".join(meta["authors"])
            bits = [b for b in [authors + "." if authors else "",
                                meta["title"] + "." if meta["title"] else "",
                                meta["journal"] + "." if meta["journal"] else "",
                                meta["vp"] + "." if meta["vp"] else "",
                                "doi: %s." % doi if doi else ""] if b]
            note = author_note(meta["authors"])
            year = int(meta["year"] or 0)
            new_entries.append({"note": note, "text": " ".join(bits), "doi": doi, "year": year})
            known_dois.add(doi.lower())
            known_titles.append(norm_title(meta["title"]))

    if new_entries:
        print("New papers auto-inserted into the CV:")
        for e in new_entries:
            print("  + (%s) %s" % (e["note"], e["text"][:80] + "..."))
        for e in new_entries:
            idx = len(pubs)
            for i, p in enumerate(pubs):
                if p.get("year", 0) < e["year"]:
                    idx = i
                    break
            pubs.insert(idx, e)
        data["publications"] = pubs
        if not args.dry_run:
            DATA_FILE.write_text(
                "# CV data - single source of truth for auto_cv.py\n"
                "# Edit this file to update the CV, then run: python auto_cv.py\n"
                + yaml.safe_dump(data, allow_unicode=True, sort_keys=False, width=10**9),
                encoding="utf-8")
            print("(cv_data.yml updated with the new entries - adjust wording/order to taste)")
    else:
        print("No new papers detected in _publications/.")

    register_fonts()
    st = styles()
    doc = BaseDocTemplate(
        str(OUT_FILE), pagesize=A4,
        leftMargin=MARGIN_L, rightMargin=MARGIN_R,
        topMargin=MARGIN_T, bottomMargin=MARGIN_B,
        title="郭祥 CV", author="Xiang Guo")
    frame = Frame(MARGIN_L, MARGIN_B, CONTENT_W, A4[1] - MARGIN_T - MARGIN_B, id="main")
    doc.addPageTemplates([PageTemplate(id="page", frames=[frame], onPage=footer)])

    if args.dry_run:
        print("dry run: nothing written (would write %s and %s)" % (OUT_FILE, XLSX_FILE))
        return

    OUT_FILE.parent.mkdir(exist_ok=True)
    doc.build(build_story(data, st))
    print("written %s (%d publications)" % (OUT_FILE, len(pubs)))

    try:
        n = build_xlsx(data, XLSX_FILE)
        print("written %s (%d achievements)" % (XLSX_FILE, n))
    except PermissionError:
        print("WARNING: %s is locked - close it in Excel/WPS and rerun to update" % XLSX_FILE)


if __name__ == "__main__":
    sys.exit(main())
